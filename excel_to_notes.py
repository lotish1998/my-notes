"""
Import book records from an Excel file into individual Markdown notes.

Usage:
    python3 excel_to_notes.py path/to/file.xlsx
"""

import sys
import os
import re
import openpyxl

OUTPUT_DIR = "notes/books"


def slugify(text):
    text = text.lower().strip()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-")


def build_note_content(row):
    title = row.get("title") or "Untitled"
    author = row.get("name.author") or "Unknown author"
    author_bio = row.get("shor.history.about.author") or ""
    description = row.get("description") or ""
    main_idea = row.get("mainIdea.fromBook") or ""

    lines = [f"# {title}", ""]
    lines.append(f"**Автор:** {author}")
    if author_bio:
        lines.append(f"\n_{author_bio}_")
    lines.append("")
    if description:
        lines.append(f"## Опис\n{description}")
        lines.append("")
    if main_idea:
        lines.append(f"## Головна ідея\n{main_idea}")

    return "\n".join(lines) + "\n"


def main():
    if len(sys.argv) < 2:
        print("Usage: python3 excel_to_notes.py path/to/file.xlsx")
        sys.exit(1)

    excel_path = sys.argv[1]
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    created = 0

    for row_cells in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, row_cells))
        if not row.get("title"):
            continue

        row_id = row.get("id", created + 1)
        slug = slugify(str(row["title"]))
        filename = f"{int(row_id):03d}-{slug}.md"
        filepath = os.path.join(OUTPUT_DIR, filename)

        content = build_note_content(row)
        with open(filepath, "w", encoding="utf-8") as f:
            f.write(content)

        created += 1

    print(f"Created {created} notes in {OUTPUT_DIR}/")


if __name__ == "__main__":
    main()
